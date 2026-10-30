from __future__ import annotations

import io
import zipfile
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.db.models import Arquivo, Empresa, Nota, Processo
from backend.app.db.session import SessionLocal, init_db
from backend.app.main import app


def _seed() -> tuple[int, int]:
    init_db()
    with SessionLocal() as db:
        db.query(Arquivo).delete()
        db.query(Nota).delete()
        db.query(Processo).delete()
        db.query(Empresa).delete()
        empresa = Empresa(nome="VS ENGENHARIA LTDA", cnpj="29003753000132", ambiente="producao", ativo=True)
        outra = Empresa(nome="OUTRA LTDA", cnpj="11222333000185", ambiente="producao", ativo=True)
        db.add_all([empresa, outra])
        db.flush()
        processo = Processo(empresa_id=empresa.id, certificado_id=1, tipo="consulta_nfse", status="finalizado")
        db.add(processo)
        db.flush()
        comum = dict(empresa_id=empresa.id, processo_id=processo.id, tomador_cnpj=empresa.cnpj, tomador_nome=empresa.nome, status_documento="autorizada")
        db.add_all(
            [
                # Fora de ordem de proposito: o layout ordena por emissao.
                Nota(
                    **comum,
                    chave="CH-2",
                    numero_nfse="751",
                    data_emissao=date(2026, 2, 19),
                    competencia=date(2026, 2, 1),
                    prestador_cnpj="30459798000103",
                    prestador_nome="PLANNING MAROX CONSULTORIA S/S LTDA",
                    valor_servico=3142.60,
                    csrf=146.13,
                    irrf=47.14,
                    valor_liquido=2949.33,
                    municipio="Goiânia-GO",
                    incidencia_iss="Goiania",
                    codigo_servico="171901",
                ),
                Nota(
                    **comum,
                    chave="CH-1",
                    numero_nfse="12",
                    data_emissao=date(2026, 1, 16),
                    competencia=date(2026, 1, 1),
                    prestador_cnpj="25197980000195",
                    prestador_nome="TRIUNFO SERVIÇOS PARA CONSTRUÇÃO CIVIL LTDA",
                    valor_servico=650,
                    inss=25.03,
                    iss_retido=True,
                    valor_iss_retido=32.50,
                    valor_liquido=592.47,
                    municipio="Rio de Janeiro",
                    incidencia_iss="Duque de Caxias",
                    codigo_servico="070201",
                ),
            ]
        )
        db.commit()
        return int(empresa.id), int(outra.id)


def _planilha(content: bytes):
    return load_workbook(io.BytesIO(content)).active


def test_planning_hub_monta_o_layout_do_modelo():
    empresa_id, _ = _seed()
    with TestClient(app) as client:
        response = client.post("/relatorios/conferencia", json={"empresa_id": empresa_id, "layout": "planning_hub"})

    assert response.status_code == 200
    assert "relacao_nfse_retencoes_vs_engenharia_ltda" in response.headers["content-disposition"]
    ws = _planilha(response.content)
    assert ws["B1"].value == "Relação de NFS-e e retenções 2026"
    assert ws["B2"].value == "Nome do Cliente: VS ENGENHARIA LTDA"
    assert ws["B3"].value == "CNPJ: 29.003.753/0001-32"
    assert ws["J6"].value == "Retenções"
    assert [ws[f"{c}7"].value for c in "JKLM"] == ["CSRF", "IRRF", "INSS", "ISS"]
    assert "J6:M6" in {str(r) for r in ws.merged_cells.ranges}

    # Linha 8 = nota mais antiga (TRIUNFO), com ISS retido e incidencia em outro municipio.
    assert ws["B8"].value == "=D8"
    assert ws["E8"].value == "25.197.980/0001-95"
    assert ws["G8"].value == 12
    assert ws["H8"].value == "07.02"
    assert ws["L8"].value == 25.03
    assert ws["M8"].value == 32.5
    assert ws["N8"].value == "=I8-J8-K8-L8-M8"
    assert ws["O8"].value == '=IF(N8<I8,"Contém Retenção","Sem Retenção")'
    assert ws["P8"].value == "Duque de Caxias"
    assert ws["R8"].value is None

    # Linha 9: mesmo municipio de incidencia (sem acento) vira "-".
    assert ws["F9"].value == "PLANNING MAROX CONSULTORIA S/S LTDA"
    assert ws["J9"].value == 146.13
    assert ws["P9"].value == "-"

    assert ws["I11"].value == "=SUBTOTAL(9,I8:I9)"
    assert ws["B13"].value.startswith("Emitido por: PLANNING MAROX")
    assert [ws[f"I{r}"].value for r in (13, 14, 15)] == ["CSRF 5952", "IRRF 1708", "INSS 1162"]
    assert ws["J13"].value == "=J11"
    assert ws.auto_filter.ref == "B7:R9"


def test_planning_hub_exige_empresa_no_filtro():
    _seed()
    with TestClient(app) as client:
        response = client.post("/relatorios/conferencia", json={"layout": "planning_hub"})

    assert response.status_code == 400
    assert "Escolha uma empresa" in response.json()["detail"]


def test_planning_ma_continua_sendo_o_padrao():
    empresa_id, _ = _seed()
    with TestClient(app) as client:
        response = client.post("/relatorios/conferencia", json={"empresa_id": empresa_id})

    assert response.status_code == 200
    abas = load_workbook(io.BytesIO(response.content)).sheetnames
    assert abas == ["Todas as Notas", "Notas Divergentes", "Notas Corretas"]


def test_layout_desconhecido_e_recusado():
    empresa_id, _ = _seed()
    with TestClient(app) as client:
        response = client.post("/relatorios/conferencia", json={"empresa_id": empresa_id, "layout": "xpto"})

    assert response.status_code == 400


def test_zip_com_planilha_planning_hub():
    empresa_id, _ = _seed()
    with TestClient(app) as client:
        response = client.get(
            "/notas/download-lote",
            params={
                "empresa_id": empresa_id,
                "incluir_xml": False,
                "incluir_pdf": True,
                "incluir_relatorio": True,
                "relatorio_layout": "planning_hub",
            },
        )

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        nomes = [n for n in zf.namelist() if n.endswith(".xlsx")]
        assert len(nomes) == 1 and "relacao_nfse_retencoes_" in nomes[0]
        ws = _planilha(zf.read(nomes[0]))
    assert ws["B2"].value == "Nome do Cliente: VS ENGENHARIA LTDA"


def test_codigo_ibge_no_lugar_do_nome_vira_o_nome_do_municipio():
    from backend.app.services.relatorio_layouts_service import _linha_hub, _nome_municipio

    assert _nome_municipio("5221403") == "Trindade"
    assert _nome_municipio("Goiânia") == "Goiânia"
    assert _nome_municipio("9999999") == "9999999"
    valores = {
        "Município": "5221403",
        "Incidência do ISS": "Trindade",
        "N° Documento": "906",
        "ISS RET. – SERV. TOMADOS": 0,
        "ISS RET. – SERV. PRESTADOS": 0,
        "Data de Emissão": date(2026, 4, 6),
        "CNPJ/CPF": "21599981000105",
        "Razão Social": "ROSENDO",
        "Código de serviço": "14.01",
        "Valor Total": 3600,
        "CSRF": 0,
        "IRRF": 0,
        "INSS": 0,
        "Data do pagamento": None,
    }
    linha = _linha_hub(valores)
    assert linha[0] == "Trindade"
    assert linha[11] == "-"
