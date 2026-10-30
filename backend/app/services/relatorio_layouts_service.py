"""Layouts da planilha de notas.

- ``planning_ma``: relatorio de conferencia de sempre (3 abas, 31 colunas), em
  ``portal_support_service.exportar_conferencia_xlsx``.
- ``planning_hub``: "Relacao de NFS-e e retencoes", copia do modelo que o
  escritorio entrega ao cliente (VS Engenharia, out/2026). Uma empresa por
  planilha: sem empresa no filtro, recusa em vez de juntar o grupo inteiro.
"""

from __future__ import annotations

import os
import re
import tempfile
import unicodedata
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from backend.app.core.tempo import agora_local, hoje_local
from backend.app.db.models import Empresa
from backend.app.schemas.notas import NotasDownloadFiltros
from backend.app.services import portal_support_service as relatorio
from backend.app.services.portal_support_service import PortalSupportError

LAYOUT_PADRAO = "planning_ma"
LAYOUTS = {"planning_ma", "planning_hub"}

# Rodape fixo do modelo: o escritorio que prepara a planilha (decisao de 20/10/2026).
EMITIDO_POR = "PLANNING MAROX CONSULTORIA S/S LTDA CNPJ: 30.459.798/0001-03"

_FONTE = "Century Gothic"
_VERDE = PatternFill("solid", fgColor="00FF99")
_FINA = Side(style="hair")
_BORDA = Border(left=_FINA, right=_FINA, top=_FINA, bottom=_FINA)
_CONTABIL = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
_DATA = "dd/mm/yyyy"
_MES = "mmm-yy"
_PRIMEIRA_LINHA = 8

# (coluna, cabecalho, largura). J..M ficam sob "Retencoes" na linha 6.
_COLUNAS = [
    ("B", "Competência", 12.7),
    ("C", "Município", 20.7),
    ("D", "Data de Emissão", 14.9),
    ("E", "CNPJ/CPF", 19.4),
    ("F", "Razão Social", 58.3),
    ("G", "Nº Documento", 13.4),
    ("H", "Cód. Serviço", 12.1),
    ("I", "Valor Total", 23.4),
    ("J", "CSRF", 13.7),
    ("K", "IRRF", 12.9),
    ("L", "INSS", 13.0),
    ("M", "ISS", 11.6),
    ("N", "Valor Líquido", 16.3),
    ("O", "Retenção", 18.7),
    ("P", "Incidência do ISS", 18.0),
    ("Q", "Data do pagamento", 18.1),
    ("R", "CNO", 30.0),
]
_DARF = [("CSRF 5952", "J"), ("IRRF 1708", "K"), ("INSS 1162", "L")]


def exportar_relatorio_xlsx(
    db: Session,
    filtros: NotasDownloadFiltros,
    grupo: str | None = None,
    layout: str | None = None,
) -> tuple[Path, str]:
    layout = layout or LAYOUT_PADRAO
    if layout not in LAYOUTS:
        raise PortalSupportError(f"Layout de planilha desconhecido: {layout}.")
    if layout == "planning_hub":
        return _exportar_planning_hub(db, filtros, grupo)
    return relatorio.exportar_conferencia_xlsx(db, filtros, grupo=grupo)


def _cnpj_mascarado(digits: str) -> str:
    if len(digits) == 14:
        return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"
    if len(digits) == 11:
        return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"
    return digits


_MUNICIPIOS_PHP = Path(__file__).resolve().parents[2] / "danfse" / "src" / "Danfse" / "Data" / "Municipios.php"


@lru_cache(maxsize=1)
def _municipios_ibge() -> dict[str, str]:
    # Mesma tabela IBGE do gerador de DANFSe (vai na imagem pelo COPY . .).
    try:
        texto = _MUNICIPIOS_PHP.read_text(encoding="utf-8")
    except OSError:
        return {}
    return dict(re.findall(r"(\d{7}) => \['nome' => '([^']+)'", texto))


def _nome_municipio(valor: str) -> str:
    """Algumas prefeituras mandam o codigo IBGE em xLocPrestacao/xLocIncid."""
    valor = (valor or "").strip()
    if len(valor) == 7 and valor.isdigit():
        return _municipios_ibge().get(valor, valor)
    return valor


def _mesmo_municipio(a: str, b: str) -> bool:
    # "Goiânia-GO" e "Goiania" sao o mesmo municipio.
    def chave(texto: str) -> str:
        sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
        return re.sub(r"[^a-z]", "", sem_acento.lower().split("-")[0])

    return bool(a) and chave(a) == chave(b)


def _empresa_do_filtro(db: Session, filtros: NotasDownloadFiltros, grupo: str | None) -> Empresa:
    if not filtros.empresa_id:
        raise PortalSupportError("Escolha uma empresa no filtro para usar o layout Planning HUB.")
    empresa = db.get(Empresa, int(filtros.empresa_id))
    if empresa is None or (grupo is not None and empresa.grupo != grupo):
        raise PortalSupportError("Empresa do filtro não encontrada.")
    return empresa


def _linha_hub(valores: dict) -> list:
    municipio = _nome_municipio(valores["Município"])
    incidencia = _nome_municipio(valores["Incidência do ISS"])
    numero = valores["N° Documento"]
    iss = valores["ISS RET. – SERV. TOMADOS"] or valores["ISS RET. – SERV. PRESTADOS"]
    return [
        municipio,
        valores["Data de Emissão"],
        _cnpj_mascarado(valores["CNPJ/CPF"]),
        valores["Razão Social"],
        int(numero) if numero.isdigit() else numero,
        valores["Código de serviço"],
        valores["Valor Total"],
        valores["CSRF"],
        valores["IRRF"],
        valores["INSS"],
        iss,
        "-" if not incidencia or _mesmo_municipio(incidencia, municipio) else incidencia,
        valores["Data do pagamento"] or "-",
    ]


def _exportar_planning_hub(db: Session, filtros: NotasDownloadFiltros, grupo: str | None) -> tuple[Path, str]:
    empresa = _empresa_do_filtro(db, filtros, grupo)
    notas = [relatorio.aplicar_campos_operacionais(nota) for nota in relatorio._buscar_notas_relatorio(db, filtros, grupo=grupo)]
    xml_resumos = relatorio._xml_resumos_por_nota(db, notas)
    empresa_cnpj = relatorio._relatorio_digits(empresa.cnpj)
    gerado_em = hoje_local()
    linhas = []
    for nota in notas:
        # Sem consulta a Receita: o layout nao mostra CNAE nem Simples.
        row = relatorio._relatorio_row(nota, xml_resumos.get(int(nota.id)), empresa_cnpj, {}, gerado_em)
        linhas.append(_linha_hub(dict(zip(relatorio.RELATORIO_XLSX_HEADERS, row))))
    # O modelo vinha fora de ordem; aqui sai por data de emissao.
    linhas.sort(key=lambda linha: linha[1] or date.min)

    datas = [linha[1] for linha in linhas if linha[1]]
    anos = sorted({d.year for d in datas}) or [gerado_em.year]
    referencia = _data(filtros.data_fim or filtros.competencia_fim) or (max(datas) if datas else gerado_em)

    wb = Workbook()
    ws = wb.active
    ws.title = "Relação NFS-e"
    _montar_planilha(ws, empresa, linhas, anos, referencia)

    nome = re.sub(r"[^A-Za-z0-9]+", "_", empresa.nome or "empresa").strip("_").lower()[:40]
    filename = f"relacao_nfse_retencoes_{nome}_{agora_local().strftime('%Y%m%d_%H%M%S')}.xlsx"
    descriptor, temp_name = tempfile.mkstemp(prefix="relacao_nfse_", suffix=".xlsx")
    os.close(descriptor)
    temp_path = Path(temp_name)
    try:
        wb.save(temp_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    return temp_path, filename


def _data(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _montar_planilha(ws, empresa: Empresa, linhas: list[list], anos: list[int], referencia: date) -> None:
    normal = Font(name=_FONTE, size=10)
    negrito = Font(name=_FONTE, size=10, bold=True)
    centro = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 5.7
    for coluna, _cabecalho, largura in _COLUNAS:
        ws.column_dimensions[coluna].width = largura

    periodo = str(anos[0]) if len(anos) == 1 else f"{anos[0]}–{anos[-1]}"
    ws["B1"] = f"Relação de NFS-e e retenções {periodo}"
    ws["B2"] = f"Nome do Cliente: {empresa.nome or ''}"
    ws["B3"] = f"CNPJ: {_cnpj_mascarado(relatorio._relatorio_digits(empresa.cnpj))}"
    ws["B4"] = referencia.replace(day=1)
    ws["B4"].number_format = _MES
    for celula in ("B1", "B2", "B3", "B4"):
        ws[celula].font = negrito
        ws[celula].alignment = Alignment(horizontal="left")

    # Cabecalho em duas linhas: tudo mesclado na vertical, menos "Retencoes".
    cabecalho = Font(name=_FONTE, size=9, bold=True)
    for coluna, titulo, _largura in _COLUNAS:
        if coluna in "JKLM":
            ws[f"{coluna}7"] = titulo
        else:
            ws[f"{coluna}6"] = titulo
            ws.merge_cells(f"{coluna}6:{coluna}7")
        for linha in (6, 7):
            celula = ws[f"{coluna}{linha}"]
            celula.font = cabecalho
            celula.fill = _VERDE
            celula.border = _BORDA
            celula.alignment = centro
    ws["J6"] = "Retenções"
    ws.merge_cells("J6:M6")

    retencao_font = Font(name=_FONTE, size=10, color="FF0000")
    for indice, valores in enumerate(linhas):
        r = _PRIMEIRA_LINHA + indice
        municipio, emissao, cnpj, razao, numero, servico, total, csrf, irrf, inss, iss, incidencia, pagamento = valores
        celulas = {
            "B": f"=D{r}",
            "C": municipio or "-",
            "D": emissao,
            "E": cnpj,
            "F": razao,
            "G": numero,
            "H": servico,
            "I": total,
            "J": csrf,
            "K": irrf,
            "L": inss,
            "M": iss,
            "N": f"=I{r}-J{r}-K{r}-L{r}-M{r}",
            "O": f'=IF(N{r}<I{r},"Contém Retenção","Sem Retenção")',
            "P": incidencia,
            "Q": pagamento,
            "R": None,  # CNO: fica em branco (decisao de 20/10/2026).
        }
        for coluna, valor in celulas.items():
            celula = ws[f"{coluna}{r}"]
            celula.value = valor
            celula.font = normal
            celula.border = _BORDA
        ws[f"B{r}"].number_format = _MES
        ws[f"B{r}"].alignment = Alignment(horizontal="center")
        ws[f"D{r}"].number_format = _DATA
        ws[f"Q{r}"].number_format = _DATA
        ws[f"O{r}"].font = retencao_font
        ws[f"O{r}"].alignment = Alignment(horizontal="center")
        for coluna in "IJKLMN":
            ws[f"{coluna}{r}"].number_format = _CONTABIL

    ultima = _PRIMEIRA_LINHA + max(len(linhas), 1) - 1
    ws.auto_filter.ref = f"B7:R{ultima}"

    total = ultima + 2
    for coluna in "IJKLMN":
        celula = ws[f"{coluna}{total}"]
        celula.value = f"=SUBTOTAL(9,{coluna}{_PRIMEIRA_LINHA}:{coluna}{ultima})"
        celula.font = negrito
        celula.border = _BORDA
        celula.number_format = _CONTABIL

    darf = total + 2
    ws[f"B{darf}"] = f"Emitido por: {EMITIDO_POR}"
    ws[f"B{darf}"].font = negrito
    for deslocamento, (rotulo, coluna) in enumerate(_DARF):
        linha = darf + deslocamento
        ws[f"I{linha}"] = rotulo
        ws[f"I{linha}"].font = negrito
        ws[f"I{linha}"].fill = _VERDE
        ws[f"J{linha}"] = f"={coluna}{total}"
        ws[f"J{linha}"].number_format = _CONTABIL
    soma = darf + len(_DARF)
    ws[f"J{soma}"] = f"=SUM(J{darf}:J{soma - 1})"
    ws[f"J{soma}"].font = negrito
    ws[f"J{soma}"].number_format = _CONTABIL
